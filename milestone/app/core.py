from __future__ import annotations

import json
import logging
import math
import os
import re
import sqlite3
import sys
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal, Protocol, TypedDict
from uuid import uuid4

import httpx
import pandas as pd
from docx import Document
from langgraph.graph import END, START, StateGraph
from langsmith import traceable
from pydantic import BaseModel, ConfigDict, Field
from sqlglot import exp, parse


ApprovalValue = Literal["PASS", "FAIL"]
WorkflowStatus = Literal["running", "approved", "rejected"]


class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid")


class SearchHit(StrictModel):
    title: str
    url: str
    content: str
    score: float = Field(ge=0)


class WebResearchInput(StrictModel):
    query: str = Field(min_length=3)
    max_results: int = Field(default=5, ge=1, le=10)


class WebResearchOutput(StrictModel):
    query: str
    provider: Literal["tavily", "local_knowledge_base"]
    results: list[SearchHit]


class TextToSQLInput(StrictModel):
    question: str = Field(min_length=3)
    candidate_sql: str | None = None


class SQLValidationResult(StrictModel):
    valid: bool
    normalized_sql: str | None = None
    errors: list[str] = Field(default_factory=list)


class TextToSQLOutput(StrictModel):
    question: str
    sql: str
    columns: list[str]
    rows: list[dict[str, Any]]
    validation_attempts: list[SQLValidationResult]
    repaired: bool


class MemoryWriteInput(StrictModel):
    thread_id: str = Field(min_length=1)
    kind: Literal["episodic", "semantic"]
    text: str = Field(min_length=1)
    metadata: dict[str, Any] = Field(default_factory=dict)


class MemorySearchInput(StrictModel):
    query: str = Field(min_length=1)
    thread_id: str | None = None
    kinds: list[Literal["episodic", "semantic"]] = Field(
        default_factory=lambda: ["episodic", "semantic"]
    )
    limit: int = Field(default=5, ge=1, le=20)


class MemoryRecord(StrictModel):
    record_id: str
    thread_id: str
    kind: Literal["episodic", "semantic"]
    text: str
    metadata: dict[str, Any]
    created_at: str
    score: float = 0.0


class MemorySearchOutput(StrictModel):
    query: str
    records: list[MemoryRecord]


class SupervisorPlan(StrictModel):
    objective: str
    steps: list[str]
    next_agent: Literal["researcher", "writer"]
    rationale: str


class ResearchBundle(StrictModel):
    web: WebResearchOutput | None = None
    sql: TextToSQLOutput | None = None
    tools_used: list[str] = Field(default_factory=list)


class ReportWriterInput(StrictModel):
    query: str
    research: ResearchBundle
    memories: list[MemoryRecord] = Field(default_factory=list)


class ReportWriterOutput(StrictModel):
    title: str
    report: str
    source_count: int = Field(ge=0)


class ApprovalRequest(StrictModel):
    request_id: str
    thread_id: str
    report_title: str
    report: str


class ApprovalDecision(StrictModel):
    decision: ApprovalValue
    reviewer: str = "queue-reviewer"
    reviewed_at: str
    source: str


class ApprovalUpdate(StrictModel):
    decision: ApprovalValue


class ApprovalQueueEntry(StrictModel):
    entry_id: str
    decision: ApprovalValue
    enqueued_at: str
    source: str


class WorkflowEvent(StrictModel):
    event: Literal["node_completed", "workflow_completed"]
    request_id: str
    node: str
    status: WorkflowStatus
    payload: dict[str, Any] = Field(default_factory=dict)


class ResearchResponse(StrictModel):
    request_id: str
    thread_id: str
    query: str
    status: Literal["approved", "rejected"]
    route_history: list[str]
    tools_used: list[str]
    approval: ApprovalDecision
    report: str | None


class ResearchGraphState(TypedDict, total=False):
    request_id: str
    thread_id: str
    query: str
    status: WorkflowStatus
    plan: dict[str, Any]
    next_agent: Literal["researcher", "writer"]
    route_history: list[str]
    memories: list[dict[str, Any]]
    research: dict[str, Any]
    report_title: str
    draft_report: str
    tools_used: list[str]
    approval: dict[str, Any]
    final_report: str | None


class JsonFormatter(logging.Formatter):
    def format(self, record: logging.LogRecord) -> str:
        payload = {
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "level": record.levelname,
            "logger": record.name,
            "message": record.getMessage(),
            "event": getattr(record, "event", "log"),
            "details": getattr(record, "details", {}),
        }
        return json.dumps(payload, default=str)


def configure_json_logging(
    log_path: Path | None = None,
    logger_name: str = "milestone4",
) -> logging.Logger:
    logger = logging.getLogger(logger_name)
    logger.setLevel(os.getenv("MILESTONE4_LOG_LEVEL", "INFO").upper())
    logger.propagate = False

    if not logger.handlers:
        console_handler = logging.StreamHandler(sys.stderr)
        console_handler.setFormatter(JsonFormatter())
        logger.addHandler(console_handler)

    if log_path is not None:
        resolved_log_path = log_path.resolve()
        has_file_handler = any(
            isinstance(handler, logging.FileHandler)
            and Path(handler.baseFilename) == resolved_log_path
            for handler in logger.handlers
        )
        if not has_file_handler:
            resolved_log_path.parent.mkdir(parents=True, exist_ok=True)
            file_handler = logging.FileHandler(resolved_log_path, encoding="utf-8")
            file_handler.setFormatter(JsonFormatter())
            logger.addHandler(file_handler)

    return logger


def log_event(logger: logging.Logger, event: str, **details: Any) -> None:
    logger.info(event, extra={"event": event, "details": details})


def _tokens(text: str) -> set[str]:
    return set(re.findall(r"[a-z0-9]+", text.lower()))


def _semantic_score(query: str, text: str) -> float:
    query_tokens = _tokens(query)
    text_tokens = _tokens(text)
    if not query_tokens or not text_tokens:
        return 0.0
    return len(query_tokens & text_tokens) / math.sqrt(
        len(query_tokens) * len(text_tokens)
    )


def _resolve_workspace_data_path(workspace_root: Path, relative_path: str) -> Path:
    """Resolve data files from the scaffolded data folder, then legacy root."""
    candidates = [
        workspace_root / "data" / relative_path,
        workspace_root / relative_path,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0]


def bootstrap_enterprise_database(
    workspace_root: Path,
    database_path: Path,
    refresh: bool = False,
) -> Path:
    workspace_root = workspace_root.resolve()
    database_path = database_path.resolve()
    if database_path.exists() and not refresh:
        return database_path

    sales_path = _resolve_workspace_data_path(workspace_root, "sales_data.csv")
    orders_path = _resolve_workspace_data_path(workspace_root, "orders.csv")
    if not sales_path.exists() or not orders_path.exists():
        raise FileNotFoundError(
            "sales_data.csv and orders.csv are required to build the SQLite database."
        )

    database_path.parent.mkdir(parents=True, exist_ok=True)
    sales_frame = pd.read_csv(sales_path)
    orders_frame = pd.read_csv(orders_path)
    with sqlite3.connect(database_path) as connection:
        sales_frame.to_sql("sales", connection, if_exists="replace", index=False)
        orders_frame.to_sql("orders", connection, if_exists="replace", index=False)
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_sales_product ON sales(Product)"
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_orders_id ON orders(Order_ID)"
        )
    return database_path


class WebResearchTool:
    def __init__(self, workspace_root: Path, logger: logging.Logger):
        self.workspace_root = workspace_root.resolve()
        self.logger = logger
        self.local_sources = self._load_local_sources()

    def _load_local_sources(self) -> list[SearchHit]:
        sources: list[SearchHit] = []
        knowledge_base_path = _resolve_workspace_data_path(
            self.workspace_root,
            "knowledge_base",
        )
        for document_path in sorted(knowledge_base_path.glob("*.docx")):
            document = Document(str(document_path))
            paragraphs = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]
            for paragraph_index, paragraph in enumerate(paragraphs):
                sources.append(
                    SearchHit(
                        title=f"{document_path.stem} section {paragraph_index + 1}",
                        url=str(document_path),
                        content=paragraph,
                        score=0.0,
                    )
                )
        return sources

    @traceable(name="web_research_tool", run_type="tool")
    def search(self, request: WebResearchInput) -> WebResearchOutput:
        tavily_api_key = os.getenv("TAVILY_API_KEY")
        if tavily_api_key:
            try:
                response = httpx.post(
                    "https://api.tavily.com/search",
                    json={
                        "api_key": tavily_api_key,
                        "query": request.query,
                        "max_results": request.max_results,
                        "search_depth": "advanced",
                    },
                    timeout=20.0,
                )
                response.raise_for_status()
                results = [
                    SearchHit(
                        title=item.get("title", "Untitled result"),
                        url=item.get("url", ""),
                        content=item.get("content", ""),
                        score=max(float(item.get("score", 0.0)), 0.0),
                    )
                    for item in response.json().get("results", [])
                ]
                log_event(
                    self.logger,
                    "tool_call",
                    tool="web_research",
                    provider="tavily",
                    result_count=len(results),
                )
                return WebResearchOutput(
                    query=request.query,
                    provider="tavily",
                    results=results,
                )
            except (httpx.HTTPError, ValueError, TypeError) as error:
                log_event(
                    self.logger,
                    "tool_fallback",
                    tool="web_research",
                    reason=str(error),
                )

        ranked_sources = [
            source.model_copy(
                update={"score": _semantic_score(request.query, source.content)}
            )
            for source in self.local_sources
        ]
        ranked_sources.sort(key=lambda source: source.score, reverse=True)
        results = ranked_sources[: request.max_results]
        log_event(
            self.logger,
            "tool_call",
            tool="web_research",
            provider="local_knowledge_base",
            result_count=len(results),
        )
        return WebResearchOutput(
            query=request.query,
            provider="local_knowledge_base",
            results=results,
        )


_FORBIDDEN_SQL_KEYS = {
    "alter",
    "attach",
    "command",
    "create",
    "delete",
    "drop",
    "insert",
    "merge",
    "pragma",
    "transaction",
    "truncate",
    "update",
}


def validate_readonly_select_sql(
    sql: str,
    allowed_tables: set[str],
    max_limit: int = 100,
) -> SQLValidationResult:
    """Shared secure-SQL guardrail: single read-only SELECT, table allowlist, capped LIMIT.

    Reused by the Milestone-3 enterprise Text-to-SQL tool and the Milestone-4
    governance Text-to-SQL tool so both enforce identical injection defenses.
    """
    errors: list[str] = []
    try:
        expressions = parse(sql, read="sqlite")
    except Exception as error:
        return SQLValidationResult(valid=False, errors=[f"SQL parse error: {error}"])

    if len(expressions) != 1:
        return SQLValidationResult(
            valid=False,
            errors=["Exactly one SQL statement is allowed."],
        )

    expression = expressions[0]
    if expression.key not in {"select", "union", "intersect", "except"}:
        errors.append("Only read-only SELECT queries are allowed.")

    if any(node.key in _FORBIDDEN_SQL_KEYS for node in expression.walk()):
        errors.append("The query contains a forbidden SQL operation.")

    referenced_tables = {
        table.name.lower() for table in expression.find_all(exp.Table)
    }
    unknown_tables = referenced_tables - allowed_tables
    if unknown_tables:
        errors.append(f"Unknown table(s): {', '.join(sorted(unknown_tables))}.")

    if errors:
        return SQLValidationResult(valid=False, errors=errors)

    limit_clause = expression.args.get("limit")
    if limit_clause is None:
        expression = expression.limit(max_limit)
    else:
        limit_expression = limit_clause.expression
        if not isinstance(limit_expression, exp.Literal) or not limit_expression.is_int:
            errors.append(f"LIMIT must be an integer no greater than {max_limit}.")
        elif int(limit_expression.this) > max_limit:
            expression = expression.limit(max_limit, copy=False)

    if errors:
        return SQLValidationResult(valid=False, errors=errors)

    return SQLValidationResult(
        valid=True,
        normalized_sql=expression.sql(dialect="sqlite"),
    )


class TextToSQLTool:
    def __init__(
        self,
        database_path: Path,
        logger: logging.Logger,
        use_llm: bool = False,
    ):
        self.database_path = database_path.resolve()
        self.logger = logger
        self.allowed_tables = {"sales", "orders"}
        self.llm = None
        if use_llm and os.getenv("OPENAI_API_KEY"):
            from langchain_openai import ChatOpenAI

            self.llm = ChatOpenAI(
                model=os.getenv("OPENAI_MODEL", "gpt-4o-mini"),
                temperature=0,
            )
        with self._connect_read_only() as connection:
            self.products = [
                row[0]
                for row in connection.execute(
                    "SELECT DISTINCT Product FROM sales ORDER BY Product"
                ).fetchall()
            ]

    def _connect_read_only(self) -> sqlite3.Connection:
        database_uri = f"file:{self.database_path.as_posix()}?mode=ro"
        connection = sqlite3.connect(database_uri, uri=True)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA query_only = ON")
        return connection

    @staticmethod
    def _requested_limit(question: str, default: int = 5) -> int:
        number_words = {
            "one": 1,
            "two": 2,
            "three": 3,
            "four": 4,
            "five": 5,
            "six": 6,
            "seven": 7,
            "eight": 8,
            "nine": 9,
            "ten": 10,
        }
        count_pattern = (
            r"\b(?:top\s+|bottom\s+|which\s+)?"
            r"(one|two|three|four|five|six|seven|eight|nine|ten|\d{1,3})"
            r"\s+(?:products?|regions?|categories?|months?|orders?|results?)\b"
        )
        match = re.search(count_pattern, question.lower())
        if not match:
            return default
        raw_count = match.group(1)
        count = number_words.get(raw_count, int(raw_count) if raw_count.isdigit() else default)
        return max(1, min(count, 100))

    def _generate_sql(self, question: str) -> str:
        if self.llm is not None:
            schema = (
                "sales(Date, Product, Category, Region, Units_Sold, Revenue, Profit); "
                "orders(Order_ID, Customer_Name, Email, Product, Quantity, Order_Date, "
                "Delivery_Date, Status, Payment_Method, Amount)"
            )
            prompt = (
                "Return one read-only SQLite SELECT query for the question. "
                "Use only the supplied schema and include LIMIT 100 or less.\n"
                f"Schema: {schema}\nQuestion: {question}"
            )
            response = self.llm.invoke(prompt)
            response_text = str(response.content).strip()
            return re.sub(r"^```(?:sql)?|```$", "", response_text).strip()

        lowered_question = question.lower()
        requested_limit = self._requested_limit(lowered_question)
        order_id_match = re.search(r"\border\s*(?:id\s*)?#?\s*(\d{4,})\b", lowered_question)
        if order_id_match:
            order_id = int(order_id_match.group(1))
            return (
                "SELECT Order_ID, Customer_Name, Product, Quantity, Order_Date, "
                "Delivery_Date, Status, Payment_Method, Amount "
                f"FROM orders WHERE Order_ID = {order_id} LIMIT 100"
            )

        if any(
            keyword in lowered_question
            for keyword in ("order", "delivery", "payment method", "shipped", "processing")
        ):
            if "status" in lowered_question or "delivered" in lowered_question or "processing" in lowered_question:
                return (
                    "SELECT Status, COUNT(*) AS Order_Count, ROUND(SUM(Amount), 2) AS Total_Amount "
                    "FROM orders GROUP BY Status ORDER BY Order_Count DESC LIMIT 100"
                )
            return (
                "SELECT Order_ID, Product, Status, Delivery_Date, Payment_Method, Amount "
                "FROM orders ORDER BY Order_Date DESC LIMIT 20"
            )

        metric = "Revenue"
        if "profit" in lowered_question:
            metric = "Profit"
        elif "unit" in lowered_question or "volume" in lowered_question:
            metric = "Units_Sold"

        direction = "ASC" if any(
            word in lowered_question for word in ("lowest", "least", "bottom")
        ) else "DESC"

        mentioned_product = next(
            (
                product
                for product in self.products
                if product.lower() in lowered_question
            ),
            None,
        )
        if mentioned_product:
            escaped_product = mentioned_product.replace("'", "''")
            return (
                "SELECT Product, ROUND(SUM(Revenue), 2) AS Total_Revenue, "
                "ROUND(SUM(Profit), 2) AS Total_Profit, SUM(Units_Sold) AS Total_Units "
                f"FROM sales WHERE Product = '{escaped_product}' GROUP BY Product LIMIT 100"
            )

        if "month" in lowered_question or "monthly" in lowered_question:
            dimension_expression = "substr(Date, 1, 7)"
            dimension_alias = "Month"
        elif "region" in lowered_question:
            dimension_expression = "Region"
            dimension_alias = "Region"
        elif "category" in lowered_question:
            dimension_expression = "Category"
            dimension_alias = "Category"
        else:
            dimension_expression = "Product"
            dimension_alias = "Product"

        return (
            f'SELECT {dimension_expression} AS "{dimension_alias}", '
            f'ROUND(SUM({metric}), 2) AS "Total_{metric}" '
            f"FROM sales GROUP BY {dimension_expression} "
            f'ORDER BY "Total_{metric}" {direction} LIMIT {requested_limit}'
        )

    def validate_sql(self, sql: str) -> SQLValidationResult:
        return validate_readonly_select_sql(sql, self.allowed_tables)

    @traceable(name="text_to_sql_tool", run_type="tool")
    def run(self, request: TextToSQLInput) -> TextToSQLOutput:
        candidate = request.candidate_sql or self._generate_sql(request.question)
        attempts: list[SQLValidationResult] = []
        repaired = False

        for attempt_number in range(2):
            validation = self.validate_sql(candidate)
            attempts.append(validation)
            if validation.valid and validation.normalized_sql:
                with self._connect_read_only() as connection:
                    cursor = connection.execute(validation.normalized_sql)
                    rows = [dict(row) for row in cursor.fetchall()]
                    columns = [description[0] for description in cursor.description or []]
                log_event(
                    self.logger,
                    "tool_call",
                    tool="text_to_sql",
                    sql=validation.normalized_sql,
                    row_count=len(rows),
                    repaired=repaired,
                )
                return TextToSQLOutput(
                    question=request.question,
                    sql=validation.normalized_sql,
                    columns=columns,
                    rows=rows,
                    validation_attempts=attempts,
                    repaired=repaired,
                )

            if attempt_number == 0:
                candidate = self._generate_sql(request.question)
                repaired = True

        error_messages = [
            error
            for validation in attempts
            for error in validation.errors
        ]
        raise ValueError("SQL validation failed: " + "; ".join(error_messages))


class MemoryStore(Protocol):
    def store(self, request: MemoryWriteInput) -> MemoryRecord: ...

    def recall(self, request: MemorySearchInput) -> MemorySearchOutput: ...


class SQLiteMemoryStore:
    def __init__(self, database_path: Path):
        self.database_path = database_path.resolve()
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        with self._connect() as connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS memory_records (
                    record_id TEXT PRIMARY KEY,
                    thread_id TEXT NOT NULL,
                    kind TEXT NOT NULL CHECK(kind IN ('episodic', 'semantic')),
                    text TEXT NOT NULL,
                    metadata_json TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )
                """
            )
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_memory_thread ON memory_records(thread_id)"
            )

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database_path, check_same_thread=False)
        connection.row_factory = sqlite3.Row
        return connection

    @traceable(name="memory_store", run_type="tool")
    def store(self, request: MemoryWriteInput) -> MemoryRecord:
        record = MemoryRecord(
            record_id=str(uuid4()),
            thread_id=request.thread_id,
            kind=request.kind,
            text=request.text,
            metadata=request.metadata,
            created_at=datetime.now(timezone.utc).isoformat(),
        )
        with self.lock, self._connect() as connection:
            connection.execute(
                """
                INSERT INTO memory_records
                    (record_id, thread_id, kind, text, metadata_json, created_at)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    record.record_id,
                    record.thread_id,
                    record.kind,
                    record.text,
                    json.dumps(record.metadata, default=str),
                    record.created_at,
                ),
            )
        return record

    @traceable(name="memory_recall", run_type="retriever")
    def recall(self, request: MemorySearchInput) -> MemorySearchOutput:
        placeholders = ",".join("?" for _ in request.kinds)
        parameters: list[Any] = list(request.kinds)
        where_parts = [f"kind IN ({placeholders})"]
        if request.thread_id:
            where_parts.append("thread_id = ?")
            parameters.append(request.thread_id)

        sql = (
            "SELECT record_id, thread_id, kind, text, metadata_json, created_at "
            "FROM memory_records WHERE "
            + " AND ".join(where_parts)
            + " ORDER BY created_at DESC LIMIT 250"
        )
        with self.lock, self._connect() as connection:
            rows = connection.execute(sql, parameters).fetchall()

        records = [
            MemoryRecord(
                record_id=row["record_id"],
                thread_id=row["thread_id"],
                kind=row["kind"],
                text=row["text"],
                metadata=json.loads(row["metadata_json"]),
                created_at=row["created_at"],
                score=_semantic_score(request.query, row["text"]),
            )
            for row in rows
        ]
        records.sort(
            key=lambda record: (record.score, record.created_at),
            reverse=True,
        )
        return MemorySearchOutput(
            query=request.query,
            records=records[: request.limit],
        )


class RedisMemoryStore:
    def __init__(self, redis_url: str):
        import redis

        self.client = redis.Redis.from_url(redis_url, decode_responses=True)
        self.client.ping()
        self.prefix = "milestone4:memory"

    @traceable(name="redis_memory_store", run_type="tool")
    def store(self, request: MemoryWriteInput) -> MemoryRecord:
        record = MemoryRecord(
            record_id=str(uuid4()),
            thread_id=request.thread_id,
            kind=request.kind,
            text=request.text,
            metadata=request.metadata,
            created_at=datetime.now(timezone.utc).isoformat(),
        )
        key = f"{self.prefix}:{request.thread_id}"
        self.client.rpush(key, record.model_dump_json())
        self.client.sadd(f"{self.prefix}:threads", request.thread_id)
        return record

    @traceable(name="redis_memory_recall", run_type="retriever")
    def recall(self, request: MemorySearchInput) -> MemorySearchOutput:
        thread_ids = (
            [request.thread_id]
            if request.thread_id
            else sorted(self.client.smembers(f"{self.prefix}:threads"))
        )
        records: list[MemoryRecord] = []
        for thread_id in thread_ids:
            if thread_id is None:
                continue
            for raw_record in self.client.lrange(
                f"{self.prefix}:{thread_id}", -250, -1
            ):
                record = MemoryRecord.model_validate_json(raw_record)
                if record.kind in request.kinds:
                    records.append(
                        record.model_copy(
                            update={
                                "score": _semantic_score(request.query, record.text)
                            }
                        )
                    )
        records.sort(
            key=lambda record: (record.score, record.created_at),
            reverse=True,
        )
        return MemorySearchOutput(
            query=request.query,
            records=records[: request.limit],
        )


def build_memory_store(
    sqlite_path: Path,
    logger: logging.Logger,
) -> MemoryStore:
    redis_url = os.getenv("REDIS_URL")
    if redis_url:
        try:
            store = RedisMemoryStore(redis_url)
            log_event(logger, "memory_backend", backend="redis")
            return store
        except Exception as error:
            log_event(
                logger,
                "memory_fallback",
                requested_backend="redis",
                fallback_backend="sqlite",
                reason=str(error),
            )
    log_event(logger, "memory_backend", backend="sqlite")
    return SQLiteMemoryStore(sqlite_path)


class ApprovalQueueGate:
    def __init__(self, queue_path: Path):
        self.queue_path = queue_path.resolve()
        self.queue_path.parent.mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        if not self.queue_path.exists():
            self._write_queue([])

    def _read_queue(self) -> list[ApprovalQueueEntry]:
        try:
            raw = self.queue_path.read_text(encoding="utf-8").strip()
            if not raw:
                return []
            payload = json.loads(raw)
        except json.JSONDecodeError as error:
            raise ValueError(
                f"Approval queue at {self.queue_path} is not valid JSON."
            ) from error

        if not isinstance(payload, list):
            raise ValueError(
                f"Approval queue at {self.queue_path} must contain a JSON array."
            )
        return [ApprovalQueueEntry.model_validate(item) for item in payload]

    def _write_queue(self, queue: list[ApprovalQueueEntry]) -> None:
        temporary_path = self.queue_path.with_suffix(self.queue_path.suffix + ".tmp")
        temporary_path.write_text(
            json.dumps([entry.model_dump() for entry in queue], indent=2) + "\n",
            encoding="utf-8",
        )
        temporary_path.replace(self.queue_path)

    def enqueue_decision(
        self,
        update: ApprovalUpdate,
        source: str = "manual",
    ) -> ApprovalQueueEntry:
        with self.lock:
            queue = self._read_queue()
            entry = ApprovalQueueEntry(
                entry_id=str(uuid4()),
                decision=update.decision,
                enqueued_at=datetime.now(timezone.utc).isoformat(),
                source=source,
            )
            queue.append(entry)
            self._write_queue(queue)
            return entry

    # Backward-compatible alias used by existing tests/callers.
    def set_decision(self, update: ApprovalUpdate) -> Path:
        self.enqueue_decision(update, source="legacy-set-decision")
        return self.queue_path

    def pending_count(self) -> int:
        with self.lock:
            return len(self._read_queue())

    @traceable(name="human_approval_tool", run_type="tool")
    def review(self, request: ApprovalRequest) -> ApprovalDecision:
        del request
        with self.lock:
            queue = self._read_queue()
            if not queue:
                raise ValueError(
                    "No approval decisions are queued. Enqueue PASS or FAIL before running research."
                )
            entry = queue.pop(0)
            self._write_queue(queue)

        return ApprovalDecision(
            decision=entry.decision,
            reviewed_at=datetime.now(timezone.utc).isoformat(),
            source=f"{self.queue_path}#{entry.entry_id}",
        )


# Compatibility alias for existing imports.
FileApprovalGate = ApprovalQueueGate


class SupervisorAgent:
    def __init__(self, logger: logging.Logger, use_llm: bool = False):
        self.logger = logger
        self.structured_llm = None
        if use_llm and os.getenv("OPENAI_API_KEY"):
            from langchain_openai import ChatOpenAI

            self.structured_llm = ChatOpenAI(
                model=os.getenv("OPENAI_MODEL", "gpt-4o-mini"),
                temperature=0,
            ).with_structured_output(SupervisorPlan)

    @traceable(name="supervisor_agent", run_type="chain")
    def plan(
        self,
        query: str,
        memories: list[MemoryRecord],
    ) -> SupervisorPlan:
        if self.structured_llm is not None:
            return self.structured_llm.invoke(
                "Plan an enterprise research task. Route to researcher unless the user "
                "only asks to recall an earlier report already present in memory. Always "
                "plan report writing and human approval.\n"
                f"Query: {query}\nMemories: {[record.text for record in memories]}"
            )

        memory_only_request = any(
            phrase in query.lower()
            for phrase in (
                "previous report",
                "earlier report",
                "what did we discuss",
                "recall our",
                "remember our",
            )
        ) and bool(memories)
        next_agent: Literal["researcher", "writer"] = (
            "writer" if memory_only_request else "researcher"
        )
        plan = SupervisorPlan(
            objective=query,
            steps=(
                ["Recall persistent memory", "Compile recalled evidence", "Request approval"]
                if memory_only_request
                else [
                    "Recall persistent memory",
                    "Gather web and/or enterprise data",
                    "Write executive report",
                    "Request approval",
                ]
            ),
            next_agent=next_agent,
            rationale=(
                "The request can be answered from persistent memory."
                if memory_only_request
                else "The request needs current research or structured enterprise data."
            ),
        )
        log_event(
            self.logger,
            "routing_decision",
            agent="supervisor",
            next_agent=plan.next_agent,
            rationale=plan.rationale,
        )
        return plan


class ResearcherAgent:
    def __init__(
        self,
        web_tool: WebResearchTool,
        sql_tool: TextToSQLTool,
        logger: logging.Logger,
    ):
        self.web_tool = web_tool
        self.sql_tool = sql_tool
        self.logger = logger

    @traceable(name="researcher_agent", run_type="chain")
    def research(self, query: str) -> ResearchBundle:
        lowered_query = query.lower()
        structured_terms = (
            "sales",
            "revenue",
            "profit",
            "unit",
            "region",
            "category",
            "order",
            "delivery",
            "payment method",
            "amount",
        )
        web_terms = (
            "competitor",
            "market",
            "trend",
            "policy",
            "warranty",
            "return",
            "refund",
            "manual",
            "shipping",
        )
        needs_sql = any(term in lowered_query for term in structured_terms)
        needs_web = any(term in lowered_query for term in web_terms) or not needs_sql

        sql_output = (
            self.sql_tool.run(TextToSQLInput(question=query)) if needs_sql else None
        )
        web_output = (
            self.web_tool.search(WebResearchInput(query=query)) if needs_web else None
        )
        tools_used = []
        if sql_output is not None:
            tools_used.append("text_to_sql")
        if web_output is not None:
            tools_used.append("web_research")

        log_event(
            self.logger,
            "agent_execution",
            agent="researcher",
            tools_used=tools_used,
        )
        return ResearchBundle(
            web=web_output,
            sql=sql_output,
            tools_used=tools_used,
        )


class WriterAgent:
    def __init__(self, logger: logging.Logger, use_llm: bool = False):
        self.logger = logger
        self.structured_llm = None
        if use_llm and os.getenv("OPENAI_API_KEY"):
            from langchain_openai import ChatOpenAI

            self.structured_llm = ChatOpenAI(
                model=os.getenv("OPENAI_MODEL", "gpt-4o-mini"),
                temperature=0,
            ).with_structured_output(ReportWriterOutput)

    @traceable(name="writer_agent", run_type="chain")
    def write(self, request: ReportWriterInput) -> ReportWriterOutput:
        if self.structured_llm is not None:
            return self.structured_llm.invoke(
                "Create a concise executive competitor-analysis report with sections for "
                "Executive Summary, Key Findings, Data Evidence, Sources, and Recommendations. "
                "Do not invent facts. The report is pending human approval.\n"
                + request.model_dump_json()
            )

        web_results = request.research.web.results if request.research.web else []
        sql_rows = request.research.sql.rows if request.research.sql else []
        memory_text = [record.text for record in request.memories[:3]]

        findings: list[str] = []
        for row in sql_rows[:5]:
            findings.append("- Enterprise data: " + json.dumps(row, default=str))
        for result in web_results[:5]:
            findings.append(f"- {result.title}: {result.content}")
        if not findings and memory_text:
            findings.extend(f"- Recalled context: {text}" for text in memory_text)
        if not findings:
            findings.append("- No grounded evidence was available for this request.")

        source_lines = [
            f"- {result.title} ({result.url})" for result in web_results[:5]
        ]
        if request.research.sql is not None:
            source_lines.append(
                f"- Enterprise SQLite query: `{request.research.sql.sql}`"
            )
        if memory_text:
            source_lines.append("- Persistent conversation memory")
        if not source_lines:
            source_lines.append("- No external sources")

        report = "\n".join(
            [
                "# Executive Research Report",
                "",
                "## Executive Summary",
                f"This report addresses: {request.query}",
                "",
                "## Key Findings",
                *findings,
                "",
                "## Data Evidence",
                (
                    json.dumps(sql_rows[:10], indent=2, default=str)
                    if sql_rows
                    else "No structured-data rows were required."
                ),
                "",
                "## Sources",
                *source_lines,
                "",
                "## Recommendations",
                "1. Validate the highest-impact finding with the responsible business owner.",
                "2. Track the cited indicators over the next reporting period.",
                "3. Release this report only after the approval gate returns PASS.",
                "",
                "## Approval Status",
                "Pending human review.",
            ]
        )
        output = ReportWriterOutput(
            title="Executive Research Report",
            report=report,
            source_count=len(source_lines),
        )
        log_event(
            self.logger,
            "agent_execution",
            agent="writer",
            source_count=output.source_count,
        )
        return output


class ResearchAssistant:
    def __init__(
        self,
        workspace_root: Path | str = Path.cwd(),
        database_path: Path | str | None = None,
        memory_path: Path | str | None = None,
        approval_path: Path | str | None = None,
        log_path: Path | str | None = None,
        use_llm: bool | None = None,
        refresh_database: bool = False,
    ):
        self.workspace_root = Path(workspace_root).resolve()
        artifacts_dir = self.workspace_root / "artifacts"
        self.database_path = Path(
            database_path or artifacts_dir / "sales.db"
        ).resolve()
        self.memory_path = Path(
            memory_path or artifacts_dir / ".milestone4_memory.db"
        ).resolve()
        self.approval_path = Path(
            approval_path or artifacts_dir / "approval_queue.json"
        ).resolve()
        self.approval_queue_path = self.approval_path
        resolved_log_path = Path(log_path).resolve() if log_path else None
        self.logger = configure_json_logging(resolved_log_path)
        self.use_llm = (
            os.getenv("MILESTONE4_USE_LLM", "false").lower() == "true"
            if use_llm is None
            else use_llm
        )

        bootstrap_enterprise_database(
            self.workspace_root,
            self.database_path,
            refresh=refresh_database,
        )
        self.memory = build_memory_store(self.memory_path, self.logger)
        self.approval_gate = ApprovalQueueGate(self.approval_queue_path)
        self.web_tool = WebResearchTool(self.workspace_root, self.logger)
        self.sql_tool = TextToSQLTool(
            self.database_path,
            self.logger,
            use_llm=self.use_llm,
        )
        self.supervisor = SupervisorAgent(self.logger, use_llm=self.use_llm)
        self.researcher = ResearcherAgent(
            self.web_tool,
            self.sql_tool,
            self.logger,
        )
        self.writer = WriterAgent(self.logger, use_llm=self.use_llm)
        self.graph = self._build_graph()

    def _build_graph(self):
        builder = StateGraph(ResearchGraphState)
        builder.add_node("supervisor", self._supervisor_node)
        builder.add_node("researcher", self._researcher_node)
        builder.add_node("writer", self._writer_node)
        builder.add_node("approval", self._approval_node)
        builder.add_node("finalize", self._finalize_node)
        builder.add_node("reject", self._reject_node)
        builder.add_edge(START, "supervisor")
        builder.add_conditional_edges(
            "supervisor",
            self._route_after_supervisor,
            {"researcher": "researcher", "writer": "writer"},
        )
        builder.add_edge("researcher", "writer")
        builder.add_edge("writer", "approval")
        builder.add_conditional_edges(
            "approval",
            self._route_after_approval,
            {"approved": "finalize", "rejected": "reject"},
        )
        builder.add_edge("finalize", END)
        builder.add_edge("reject", END)
        return builder.compile()

    def _supervisor_node(self, state: ResearchGraphState) -> ResearchGraphState:
        memories = self.memory.recall(
            MemorySearchInput(
                query=state["query"],
                thread_id=state["thread_id"],
                limit=5,
            )
        )
        plan = self.supervisor.plan(state["query"], memories.records)
        return {
            "plan": plan.model_dump(),
            "next_agent": plan.next_agent,
            "memories": [record.model_dump() for record in memories.records],
            "route_history": [*state.get("route_history", []), "supervisor"],
        }

    @staticmethod
    def _route_after_supervisor(
        state: ResearchGraphState,
    ) -> Literal["researcher", "writer"]:
        return state["next_agent"]

    def _researcher_node(self, state: ResearchGraphState) -> ResearchGraphState:
        research = self.researcher.research(state["query"])
        return {
            "research": research.model_dump(),
            "tools_used": research.tools_used,
            "route_history": [*state.get("route_history", []), "researcher"],
        }

    def _writer_node(self, state: ResearchGraphState) -> ResearchGraphState:
        research = ResearchBundle.model_validate(state.get("research", {}))
        memories = [
            MemoryRecord.model_validate(record)
            for record in state.get("memories", [])
        ]
        output = self.writer.write(
            ReportWriterInput(
                query=state["query"],
                research=research,
                memories=memories,
            )
        )
        return {
            "report_title": output.title,
            "draft_report": output.report,
            "route_history": [*state.get("route_history", []), "writer"],
        }

    def _approval_node(self, state: ResearchGraphState) -> ResearchGraphState:
        decision = self.approval_gate.review(
            ApprovalRequest(
                request_id=state["request_id"],
                thread_id=state["thread_id"],
                report_title=state["report_title"],
                report=state["draft_report"],
            )
        )
        log_event(
            self.logger,
            "approval_decision",
            request_id=state["request_id"],
            decision=decision.decision,
            source=decision.source,
        )
        return {
            "approval": decision.model_dump(),
            "route_history": [*state.get("route_history", []), "approval"],
        }

    @staticmethod
    def _route_after_approval(
        state: ResearchGraphState,
    ) -> Literal["approved", "rejected"]:
        return "approved" if state["approval"]["decision"] == "PASS" else "rejected"

    def _store_completion_memory(
        self,
        state: ResearchGraphState,
        status: Literal["approved", "rejected"],
    ) -> None:
        report_text = state["draft_report"] if status == "approved" else "Report withheld"
        self.memory.store(
            MemoryWriteInput(
                thread_id=state["thread_id"],
                kind="episodic",
                text=f"Query: {state['query']}\nStatus: {status}\n{report_text}",
                metadata={
                    "request_id": state["request_id"],
                    "status": status,
                    "route_history": state["route_history"],
                },
            )
        )
        if status == "approved":
            self.memory.store(
                MemoryWriteInput(
                    thread_id=state["thread_id"],
                    kind="semantic",
                    text=state["draft_report"],
                    metadata={"request_id": state["request_id"]},
                )
            )

    def _finalize_node(self, state: ResearchGraphState) -> ResearchGraphState:
        route_history = [*state.get("route_history", []), "finalize"]
        approval = ApprovalDecision.model_validate(state["approval"])
        approved_report = state["draft_report"].replace(
            "Pending human review.",
            f"Approved by {approval.reviewer} at {approval.reviewed_at}.",
        )
        completed_state = {
            **state,
            "draft_report": approved_report,
            "route_history": route_history,
        }
        self._store_completion_memory(completed_state, "approved")
        log_event(
            self.logger,
            "workflow_completed",
            request_id=state["request_id"],
            status="approved",
        )
        return {
            "status": "approved",
            "draft_report": approved_report,
            "final_report": approved_report,
            "route_history": route_history,
        }

    def _reject_node(self, state: ResearchGraphState) -> ResearchGraphState:
        route_history = [*state.get("route_history", []), "reject"]
        completed_state = {**state, "route_history": route_history}
        self._store_completion_memory(completed_state, "rejected")
        log_event(
            self.logger,
            "workflow_completed",
            request_id=state["request_id"],
            status="rejected",
        )
        return {
            "status": "rejected",
            "final_report": None,
            "route_history": route_history,
        }

    @staticmethod
    def _initial_state(query: str, thread_id: str) -> ResearchGraphState:
        return {
            "request_id": str(uuid4()),
            "thread_id": thread_id,
            "query": query,
            "status": "running",
            "route_history": [],
            "tools_used": [],
            "research": {},
            "final_report": None,
        }

    @staticmethod
    def _response_from_state(state: ResearchGraphState) -> ResearchResponse:
        return ResearchResponse(
            request_id=state["request_id"],
            thread_id=state["thread_id"],
            query=state["query"],
            status=state["status"],
            route_history=state["route_history"],
            tools_used=state.get("tools_used", []),
            approval=ApprovalDecision.model_validate(state["approval"]),
            report=state.get("final_report"),
        )

    @traceable(name="milestone4_research_workflow", run_type="chain")
    def run(self, query: str, thread_id: str = "default") -> ResearchResponse:
        final_state = self.graph.invoke(self._initial_state(query, thread_id))
        return self._response_from_state(final_state)

    @traceable(name="milestone4_streaming_workflow", run_type="chain")
    def stream(self, query: str, thread_id: str = "default"):
        current_state = self._initial_state(query, thread_id)
        for update in self.graph.stream(current_state, stream_mode="updates"):
            for node, node_update in update.items():
                current_state.update(node_update)
                yield WorkflowEvent(
                    event="node_completed",
                    request_id=current_state["request_id"],
                    node=node,
                    status=current_state.get("status", "running"),
                    payload={
                        "route_history": current_state.get("route_history", []),
                        "tools_used": current_state.get("tools_used", []),
                        "approval": current_state.get("approval"),
                    },
                )

        response = self._response_from_state(current_state)
        yield WorkflowEvent(
            event="workflow_completed",
            request_id=current_state["request_id"],
            node="workflow",
            status=current_state["status"],
            payload={"response": response.model_dump()},
        )


__all__ = [
    "ApprovalQueueEntry",
    "ApprovalQueueGate",
    "ApprovalDecision",
    "ApprovalRequest",
    "ApprovalUpdate",
    "FileApprovalGate",
    "JsonFormatter",
    "MemoryRecord",
    "MemorySearchInput",
    "MemorySearchOutput",
    "MemoryWriteInput",
    "ResearchAssistant",
    "ResearchResponse",
    "SQLiteMemoryStore",
    "SQLValidationResult",
    "TextToSQLInput",
    "TextToSQLOutput",
    "TextToSQLTool",
    "WebResearchInput",
    "WebResearchOutput",
    "WorkflowEvent",
    "bootstrap_enterprise_database",
    "configure_json_logging",
    "validate_readonly_select_sql",
]